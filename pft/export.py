"""Excel workbook export. Every sheet is built from the same Report the UI renders."""

from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pft.data import month_name
from pft.report import Report

MONEY_FMT = '#,##0.00;-#,##0.00'
PCT_FMT = '+0.0%;-0.0%;0.0%'
HEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
HEADER_FONT = Font(bold=True, color="171717")
TITLE_FONT = Font(bold=True, size=14, color="171717")
MUTED_FONT = Font(italic=True, color="475569")
THIN = Side(style="thin", color="CBD5E1")

DISCLAIMER = ("Synthetic demo data. Fictional programs, vendors and amounts; not a production system, "
              "not audited, and not a certified internal-control or ESG reporting system.")

# (header, column key, kind). kind: money (cents -> CAD), pct, int, num, text
BVA_SPEC = [("Program", "program", "text"), ("Month", "month", "text"),
            ("Budget (CAD)", "budget_cents", "money"), ("Actual (CAD)", "actual_cents", "money"),
            ("Variance (CAD)", "variance_cents", "money"), ("Variance %", "variance_pct", "pct"),
            ("Remaining (CAD)", "remaining_cents", "money"), ("Expenses", "txn_count", "int")]
PROGRAM_SPEC = [("Program", "program", "text"), ("Budget (CAD)", "budget_cents", "money"),
                ("Actual (CAD)", "actual_cents", "money"), ("Variance (CAD)", "variance_cents", "money"),
                ("Variance %", "variance_pct", "pct"), ("Remaining (CAD)", "remaining_cents", "money"),
                ("Expenses", "txn_count", "int")]
EXC_SPEC = [("Exception ID", "exception_id", "text"), ("Rule", "rule_label", "text"),
            ("Category", "category", "text"), ("Severity", "severity", "text"),
            ("Status", "status", "text"), ("Record type", "record_type", "text"),
            ("Record ID", "record_id", "text"), ("Program", "program", "text"), ("Month", "month", "text"),
            ("Recorded (CAD)", "recorded_cents", "money"), ("Supported / approved (CAD)", "supported_cents", "money"),
            ("Difference (CAD)", "difference_cents", "money"), ("Related expenses", "related_txns", "text"),
            ("Explanation", "explanation", "text"), ("Review note", "note", "text")]
CTL_SPEC = [("Control ID", "control_id", "text"), ("Control", "name", "text"), ("Scope", "scope", "text"),
            ("Evidence reference", "evidence_ref", "text"), ("Owner", "owner", "text"),
            ("Due date", "due_date", "text"), ("Status", "status", "text"), ("Notes", "notes", "text"),
            ("Last updated (UTC)", "updated_at", "text")]
ESG_SPEC = [("Indicator", "name", "text"), ("Definition", "definition", "text"), ("Scope", "scope", "text"),
            ("Period", "period", "text"), ("Value", "value", "num"), ("Unit", "unit", "text"),
            ("Expected unit", "expected_unit", "text"), ("Source", "source", "text"), ("Owner", "owner", "text"),
            ("Verification status", "verification_status", "text"), ("Data-quality flags", "issues", "text")]

DICTIONARY = [
    ("All", "Amounts", "Stored as integer CAD cents in SQLite; shown and exported in CAD dollars with two decimals."),
    ("All", "Scope", "Program and month filter active in the app when this workbook was downloaded."),
    ("Budget vs Actual", "Budget (CAD)", "Approved monthly program budget. One budget row per program-month."),
    ("Budget vs Actual", "Actual (CAD)", "Sum of recorded expenses for the program-month, including flagged expenses."),
    ("Budget vs Actual", "Variance (CAD)", "Actual minus budget. Positive means overspend."),
    ("Budget vs Actual", "Variance %", "Variance divided by budget. N/A when budget is zero."),
    ("Budget vs Actual", "Remaining (CAD)", "Budget minus actual. Negative means the budget is exhausted."),
    ("Budget vs Actual", "Expenses", "Number of expense records in the program-month."),
    ("Reconciliation Exceptions", "Rule", "Matching rule that raised the exception."),
    ("Reconciliation Exceptions", "Category", "Financial error: amounts disagree or a receipt is reused. "
     "Control breach: unapproved vendor or spend above a PO. Missing support: receipt absent or partial; not a proven error."),
    ("Reconciliation Exceptions", "Recorded (CAD)", "Expense amount as recorded; for duplicates, total of all expenses citing the receipt; for POs, total spend on the PO."),
    ("Reconciliation Exceptions", "Supported / approved (CAD)", "Sum of matched receipts, or the approved PO amount. Blank where not applicable."),
    ("Reconciliation Exceptions", "Difference (CAD)", "Recorded minus supported/approved."),
    ("Reconciliation Exceptions", "Month / Program", "Month and program of the record that triggered the exception."),
    ("Reconciliation Exceptions", "Status", "Review status recorded in the app: Open, Under review or Resolved."),
    ("Control Evidence", "Evidence reference", "Identifier of the supporting document in the evidence folder (fictional)."),
    ("Control Evidence", "Scope", "Program the control applies to, or Organization-wide. Filtered by due-date month."),
    ("ESG Data", "Value", "Reported value in the stated unit. Blank means missing; it is never treated as zero."),
    ("ESG Data", "Data-quality flags", "Missing value, missing source, or a unit different from the indicator's expected unit. Units are not converted."),
    ("ESG Data", "Verification status", "Internal review status of the record, not an external audit or assurance opinion."),
]


def _cell_value(value, kind):
    if value is None or (isinstance(value, float) and pd.isna(value)) or value is pd.NA:
        return "N/A" if kind == "pct" else None
    if kind == "money":
        return int(value) / 100
    if kind == "int":
        return int(value)
    if kind in ("pct", "num"):
        return float(value)
    return str(value)


def _write_table(ws, df: pd.DataFrame, spec, start_row: int, autofilter: bool = True) -> int:
    for c, (header, _, kind) in enumerate(spec, start=1):
        cell = ws.cell(row=start_row, column=c, value=header)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.border = Border(bottom=THIN)
        cell.alignment = Alignment(horizontal="right" if kind in ("money", "pct", "int", "num") else "left",
                                   vertical="center", wrap_text=True)
    r = start_row
    for r, row in enumerate(df.to_dict("records"), start=start_row + 1):
        for c, (_, key, kind) in enumerate(spec, start=1):
            cell = ws.cell(row=r, column=c, value=_cell_value(row.get(key), kind))
            if kind == "money":
                cell.number_format = MONEY_FMT
            elif kind == "pct":
                cell.number_format = PCT_FMT
                cell.alignment = Alignment(horizontal="right")
            elif kind == "num":
                cell.number_format = "#,##0.##"
    if autofilter:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(spec))}{max(r, start_row)}"
    return r


def _sheet_header(ws, title: str, report: Report, generated: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Scope: {report.scope_label}   |   Generated: {generated}   |   Synthetic demo data"
    ws["A2"].font = MUTED_FONT


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _with_program_names(df: pd.DataFrame, programs: dict) -> pd.DataFrame:
    df = df.copy()
    df["program"] = df["program_id"].map(programs)
    return df


def build_workbook(report: Report, program_names: dict[str, str], generated: str | None = None) -> Workbook:
    generated = generated or datetime.now().strftime("%Y-%m-%d %H:%M")
    wb = Workbook()

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Program Finance Tracker - Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = DISCLAIMER
    ws["A2"].font = MUTED_FONT
    ws["A4"], ws["B4"] = "Programs in scope", report.programs_label
    ws["A5"], ws["B5"] = "Months in scope", report.months_label
    ws["A6"], ws["B6"] = "Generated", generated
    for r in (4, 5, 6):
        ws[f"A{r}"].font = Font(bold=True)
    t = report.totals
    kpis = [("Budget (CAD)", t["budget_cents"], "money"), ("Actual spend (CAD)", t["actual_cents"], "money"),
            ("Remaining budget (CAD)", t["remaining_cents"], "money"),
            ("Variance (CAD), actual - budget", t["variance_cents"], "money"),
            ("Variance %", t["variance_pct"], "pct"),
            ("Open exceptions", report.open_exceptions, "int")]
    ws["A8"] = "Key figures"
    ws["A8"].font = HEADER_FONT
    for i, (label, value, kind) in enumerate(kpis, start=9):
        ws.cell(row=i, column=1, value=label)
        cell = ws.cell(row=i, column=2, value=_cell_value(value, kind))
        cell.number_format = {"money": MONEY_FMT, "pct": PCT_FMT}.get(kind, "0")
        cell.alignment = Alignment(horizontal="right")
    row = 9 + len(kpis) + 1
    ws.cell(row=row, column=1, value="Program variance").font = HEADER_FONT
    end = _write_table(ws, report.programs, PROGRAM_SPEC, row + 1, autofilter=False)
    row = end + 2
    ws.cell(row=row, column=1, value="Findings").font = HEADER_FONT
    for i, text in enumerate(report.findings, start=row + 1):
        ws.cell(row=i, column=1, value=f"- {text}")
    _widths(ws, [34, 22, 18, 18, 14, 18, 12])

    # Budget vs Actual
    ws = wb.create_sheet("Budget vs Actual")
    _sheet_header(ws, "Budget vs Actual by program and month", report, generated)
    bva = report.bva.copy()
    bva["month"] = bva["month"].map(month_name)
    _write_table(ws, bva, BVA_SPEC, 4)
    ws.freeze_panes = "A5"
    _widths(ws, [28, 12, 16, 16, 16, 12, 16, 10])

    # Reconciliation Exceptions
    ws = wb.create_sheet("Reconciliation Exceptions")
    _sheet_header(ws, "Reconciliation exceptions (flagged, not corrected)", report, generated)
    exc = _with_program_names(report.exceptions, program_names)
    exc["month"] = exc["month"].map(month_name)
    _write_table(ws, exc, EXC_SPEC, 4)
    ws.freeze_panes = "B5"
    _widths(ws, [30, 24, 16, 10, 13, 14, 13, 26, 11, 16, 18, 16, 30, 70, 30])

    # Control Evidence
    ws = wb.create_sheet("Control Evidence")
    _sheet_header(ws, "Control evidence checklist (demo workflow, not a certified control system)", report, generated)
    _write_table(ws, report.controls, CTL_SPEC, 4)
    ws.freeze_panes = "B5"
    _widths(ws, [11, 36, 26, 22, 18, 12, 16, 40, 22])

    # ESG Data
    ws = wb.create_sheet("ESG Data")
    _sheet_header(ws, "ESG indicator data (synthetic, unaudited)", report, generated)
    esg = report.esg.copy()
    esg["period"] = esg["period"].map(month_name)
    _write_table(ws, esg, ESG_SPEC, 4)
    ws.freeze_panes = "B5"
    _widths(ws, [24, 50, 26, 11, 11, 13, 14, 34, 18, 18, 46])

    # Data Dictionary
    ws = wb.create_sheet("Data Dictionary")
    ws["A1"] = "Data dictionary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = DISCLAIMER
    ws["A2"].font = MUTED_FONT
    dd = pd.DataFrame(DICTIONARY, columns=["sheet", "field", "definition"])
    _write_table(ws, dd, [("Sheet", "sheet", "text"), ("Field", "field", "text"),
                          ("Definition", "definition", "text")], 4)
    ws.freeze_panes = "A5"
    _widths(ws, [26, 26, 110])
    return wb


def workbook_bytes(report: Report, program_names: dict[str, str]) -> bytes:
    buf = io.BytesIO()
    build_workbook(report, program_names).save(buf)
    return buf.getvalue()
