"""The Excel workbook must show the same numbers as the UI for the same filters."""

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from pft.data import Scope, full_scope
from pft.export import workbook_bytes
from pft.money import fmt_cad_whole
from pft.report import build_report

ROOT = Path(__file__).resolve().parent.parent
SHEETS = ["Summary", "Budget vs Actual", "Reconciliation Exceptions", "Control Evidence", "ESG Data",
          "Data Dictionary"]


def _names(frames):
    return dict(zip(frames["programs"]["program_id"], frames["programs"]["name"]))


def _workbook(frames, scope):
    report = build_report(frames, scope)
    return report, load_workbook(io.BytesIO(workbook_bytes(report, _names(frames))))


def _summary(ws):
    return {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 20)}


def _data_rows(ws):
    return [row for row in ws.iter_rows(min_row=5, values_only=True) if any(v is not None for v in row)]


@pytest.mark.parametrize("scope", [None, Scope(("P02", "P04"), ("2026-06", "2026-07", "2026-08"))])
def test_workbook_matches_report(frames, scope):
    scope = scope or full_scope(frames)
    report, wb = _workbook(frames, scope)
    assert wb.sheetnames == SHEETS

    s = _summary(wb["Summary"])
    t = report.totals
    assert round(s["Budget (CAD)"] * 100) == t["budget_cents"]
    assert round(s["Actual spend (CAD)"] * 100) == t["actual_cents"]
    assert round(s["Remaining budget (CAD)"] * 100) == t["remaining_cents"]
    assert s["Open exceptions"] == report.open_exceptions
    assert s["Months in scope"] == report.months_label

    bva = _data_rows(wb["Budget vs Actual"])
    assert len(bva) == len(report.bva)
    assert round(sum(r[2] for r in bva) * 100) == t["budget_cents"]
    assert round(sum(r[3] for r in bva) * 100) == t["actual_cents"]

    assert len(_data_rows(wb["Reconciliation Exceptions"])) == len(report.exceptions)
    assert len(_data_rows(wb["Control Evidence"])) == len(report.controls)
    assert len(_data_rows(wb["ESG Data"])) == len(report.esg)


def test_workbook_formatting(frames):
    _, wb = _workbook(frames, full_scope(frames))
    ws = wb["Budget vs Actual"]
    assert ws.freeze_panes == "A5"
    assert ws.auto_filter.ref.startswith("A4")
    assert "Scope:" in ws["A2"].value
    assert ws["C5"].number_format == "#,##0.00;-#,##0.00"
    for name in SHEETS[1:]:
        assert wb[name].freeze_panes is not None


def test_zero_budget_and_missing_esg_in_workbook(frames):
    _, wb = _workbook(frames, Scope(("P02", "P03"), ("2026-04", "2026-08")))
    rows = _data_rows(wb["Budget vs Actual"])
    zero = [r for r in rows if r[0] == "Community Literacy" and r[1] == "Apr 2026"][0]
    assert zero[2] == 0 and zero[5] == "N/A"
    esg = _data_rows(wb["ESG Data"])
    missing = [r for r in esg if r[2] == "Rural Health Access" and r[3] == "Aug 2026"][0]
    assert missing[4] is None  # blank, never 0
    assert "Missing value" in missing[10]


def test_ui_metrics_match_workbook(db_file, frames):
    """Run the real app with a filter applied and compare its KPI cards to the export."""
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
    at.run()
    assert not at.exception
    at.multiselect(key="f_programs").set_value(["P01", "P03"])
    at.select_slider(key="f_months").set_value(("2026-05", "2026-08"))
    at.run()
    assert not at.exception
    at.run()  # filters must survive a further rerun (any widget interaction triggers one)
    assert not at.exception
    assert at.select_slider(key="f_months").value == ("2026-05", "2026-08")
    ui = {m.label: m.value for m in at.metric}

    _, wb = _workbook(frames, Scope(("P01", "P03"), ("2026-05", "2026-06", "2026-07", "2026-08")))
    s = _summary(wb["Summary"])
    assert ui["Budget"] == fmt_cad_whole(round(s["Budget (CAD)"] * 100))
    assert ui["Actual spend"] == fmt_cad_whole(round(s["Actual spend (CAD)"] * 100))
    assert ui["Remaining budget"] == fmt_cad_whole(round(s["Remaining budget (CAD)"] * 100))
    assert ui["Open exceptions"] == str(s["Open exceptions"])


def test_ui_reset_and_empty_state(db_file, frames):
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
    at.run()
    at.multiselect(key="f_programs").set_value([]).run()
    assert not at.exception
    assert any("No programs selected" in i.value for i in at.info)
    assert not at.metric  # no KPI cards without a scope

    at.button(key="empty_reset").click().run()  # the in-page reset offered by the empty state
    assert not at.exception
    assert at.multiselect(key="f_programs").value == ["P01", "P02", "P03", "P04"]
    budget = {m.label: m.value for m in at.metric}["Budget"]
    assert budget == fmt_cad_whole(int(frames["budgets"]["budget_cents"].sum()))
